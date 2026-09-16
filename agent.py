"""
AI Agent: Automated Food Image Collection & Processing
Excel -> Image Search -> Select -> Resize 1800x1200 -> Rename -> Drive/Local -> Report
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import os
import re
import time
from pathlib import Path

import requests
from openpyxl import load_workbook
from PIL import Image, ImageFilter, ImageStat

TARGET_W, TARGET_H = 1800, 1200
MAX_BYTES = 10 * 1024 * 1024
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    )
}


def read_food_items(excel_path: str) -> list[str]:
    wb = load_workbook(excel_path, read_only=True)
    ws = wb.active
    items: list[str] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if not row or row[0] is None:
            continue
        name = str(row[0]).strip()
        if not name or name.lower() in {"food item", "item_name", "item name"}:
            continue
        items.append(name)
    wb.close()
    return items


def sanitize_filename(name: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*]', "", name).strip()
    return cleaned or "food_item"


def search_image_urls(query: str, max_results: int = 8) -> list[str]:
    """Search DuckDuckGo Images for food photos."""
    urls: list[str] = []
    try:
        from ddgs import DDGS
    except ImportError:
        from duckduckgo_search import DDGS  # type: ignore

    q = f'"{query}" food dish recipe restaurant plated'
    with DDGS() as ddgs:
        results = list(ddgs.images(q, max_results=max_results))
    skip = ("nebula", "galaxy", "astronomy", "wallpaper", "stock-vector")
    for r in results:
        url = (r.get("image") or r.get("url") or "").strip()
        title = (r.get("title") or "").lower()
        if not url.startswith("http"):
            continue
        if any(s in url.lower() or s in title for s in skip):
            continue
        urls.append(url)
    return urls


def download_image(url: str, timeout: int = 20) -> Image.Image | None:
    try:
        resp = requests.get(url, headers=HEADERS, timeout=timeout, stream=True)
        resp.raise_for_status()
        ctype = (resp.headers.get("Content-Type") or "").lower()
        if "image" not in ctype and not url.lower().endswith(
            (".jpg", ".jpeg", ".png", ".webp")
        ):
            return None
        data = resp.content
        if len(data) < 5000:
            return None
        img = Image.open(io.BytesIO(data))
        img = img.convert("RGB")
        return img
    except Exception:
        return None


def blur_score(img: Image.Image) -> float:
    gray = img.convert("L")
    edges = gray.filter(ImageFilter.FIND_EDGES)
    return float(ImageStat.Stat(edges).var[0])


def evaluate_image(img: Image.Image) -> tuple[bool, str]:
    w, h = img.size
    if w < 400 or h < 300:
        return False, "too_small"
    if blur_score(img) < 40:
        return False, "blurry"
    # Reject extreme crops (very narrow aspect) that look zoomed/cropped badly
    aspect = w / h
    if aspect < 0.7 or aspect > 2.2:
        return False, "bad_aspect"
    return True, "ok"


def process_image(img: Image.Image) -> bytes:
    """Center-crop to 3:2 then resize to 1800x1200, JPEG under 10MB."""
    target_aspect = TARGET_W / TARGET_H
    w, h = img.size
    aspect = w / h
    if aspect > target_aspect:
        new_w = int(h * target_aspect)
        left = (w - new_w) // 2
        img = img.crop((left, 0, left + new_w, h))
    else:
        new_h = int(w / target_aspect)
        top = (h - new_h) // 2
        img = img.crop((0, top, w, top + new_h))

    img = img.resize((TARGET_W, TARGET_H), Image.Resampling.LANCZOS)

    quality = 90
    while quality >= 50:
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=quality, optimize=True)
        data = buf.getvalue()
        if len(data) < MAX_BYTES:
            return data
        quality -= 10
    return data


def upload_to_drive(file_path: Path, folder_id: str, creds_path: str) -> str | None:
    """Upload via Google service account. Returns file web link or None."""
    try:
        from google.oauth2 import service_account
        from googleapiclient.discovery import build
        from googleapiclient.http import MediaFileUpload
    except ImportError:
        print("  Drive libs missing; skipping upload (pip install google-api-python-client google-auth)")
        return None

    if not os.path.exists(creds_path):
        print(f"  No credentials at {creds_path}; skipping Drive upload")
        return None

    scopes = ["https://www.googleapis.com/auth/drive.file"]
    creds = service_account.Credentials.from_service_account_file(creds_path, scopes=scopes)
    service = build("drive", "v3", credentials=creds)
    meta = {"name": file_path.name, "parents": [folder_id]}
    media = MediaFileUpload(str(file_path), mimetype="image/jpeg")
    f = (
        service.files()
        .create(body=meta, media_body=media, fields="id,webViewLink")
        .execute()
    )
    return f.get("webViewLink")


def process_item(
    name: str,
    out_dir: Path,
    max_candidates: int = 6,
) -> dict:
    row = {
        "Food Item": name,
        "Image Found": "No",
        "Image Processed": "No",
        "Uploaded": "No",
        "Source": "",
        "Status": "Failed",
        "Local Path": "",
        "Drive Link": "",
        "Notes": "",
    }
    try:
        urls = search_image_urls(name, max_results=max_candidates)
        if not urls:
            row["Notes"] = "no_search_results"
            return row

        chosen = None
        source = ""
        last_reason = ""
        for url in urls:
            img = download_image(url)
            if img is None:
                last_reason = "download_failed"
                continue
            ok, reason = evaluate_image(img)
            if not ok:
                last_reason = reason
                continue
            chosen = img
            source = url
            break

        if chosen is None:
            row["Notes"] = f"no_valid_image:{last_reason}"
            return row

        row["Image Found"] = "Yes"
        row["Source"] = source

        data = process_image(chosen)
        fname = f"{sanitize_filename(name)}.jpg"
        path = out_dir / fname
        path.write_bytes(data)
        row["Image Processed"] = "Yes"
        row["Local Path"] = str(path)
        row["Status"] = "Processed"
        row["Notes"] = f"size_bytes={len(data)};dims={TARGET_W}x{TARGET_H}"
        return row
    except Exception as e:
        row["Notes"] = f"error:{e}"
        return row


def write_report(rows: list[dict], report_path: Path) -> None:
    fields = [
        "Food Item",
        "Image Found",
        "Image Processed",
        "Uploaded",
        "Source",
        "Status",
        "Local Path",
        "Drive Link",
        "Notes",
    ]
    with report_path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description="Food image automation agent")
    parser.add_argument(
        "--excel",
        default="sample_food_items.xlsx",
        help="Input Excel with Food Item column",
    )
    parser.add_argument("--out", default="output_images", help="Local output folder")
    parser.add_argument("--report", default="processing_report.csv")
    parser.add_argument("--limit", type=int, default=0, help="Process only first N items (0=all)")
    parser.add_argument("--drive-folder-id", default=os.getenv("DRIVE_FOLDER_ID", ""))
    parser.add_argument(
        "--credentials",
        default=os.getenv("GOOGLE_APPLICATION_CREDENTIALS", "credentials.json"),
    )
    parser.add_argument("--delay", type=float, default=1.0, help="Delay between items")
    args = parser.parse_args()

    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)

    items = read_food_items(args.excel)
    if args.limit and args.limit > 0:
        items = items[: args.limit]

    print(f"Processing {len(items)} food items from {args.excel}")
    rows: list[dict] = []

    for i, name in enumerate(items, 1):
        print(f"[{i}/{len(items)}] {name}")
        row = process_item(name, out_dir)
        if (
            row["Image Processed"] == "Yes"
            and args.drive_folder_id
            and Path(args.credentials).exists()
        ):
            link = upload_to_drive(Path(row["Local Path"]), args.drive_folder_id, args.credentials)
            if link:
                row["Uploaded"] = "Yes"
                row["Drive Link"] = link
                row["Status"] = "Uploaded"
            else:
                row["Status"] = "Processed (Drive skipped)"
        elif row["Image Processed"] == "Yes":
            row["Status"] = "Processed (local only; add credentials.json + DRIVE_FOLDER_ID to upload)"
        rows.append(row)
        print(f"  -> {row['Status']} | {row['Notes']}")
        write_report(rows, Path(args.report))
        if i < len(items):
            time.sleep(args.delay)

    ok = sum(1 for r in rows if r["Image Processed"] == "Yes")
    print(f"\nDone. {ok}/{len(rows)} processed. Report: {args.report}")
    print(f"Images: {out_dir.resolve()}")


if __name__ == "__main__":
    main()
